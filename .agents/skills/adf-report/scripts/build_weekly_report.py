#!/usr/bin/env python3
"""Fill the ADF weekly-report workbook template from a JSON data file.

Usage: python3 build_weekly_report.py --data weekly-data.json --out DIR [--template XLSX]

The template (../assets/weekly-report-template.xlsx) is loaded as-is, so every style, width, banner and guide text
comes from it. This script only writes data, grows or shrinks each table to the number of rows supplied, adds sprint
columns beyond the template's six, and rewrites every formula against the final row positions.

The data contract is ../references/weekly-data-contract.md. Requires openpyxl (not standard library):
  uv run --with openpyxl python3 build_weekly_report.py ...   or   python3 -m pip install --user openpyxl

Exit codes: 0 = written, 1 = invalid data, 2 = openpyxl or the template is missing.
"""

from __future__ import annotations

import argparse
import datetime
import json
import math
import re
import sys
from copy import copy
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import MultiCellRange
    from openpyxl.worksheet.dimensions import ColumnDimension
except ImportError:
    print("openpyxl is required: `uv run --with openpyxl python3 ...` or `python3 -m pip install --user openpyxl`",
          file=sys.stderr)
    sys.exit(2)

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "weekly-report-template.xlsx"
FORBIDDEN_IN_NAME = re.compile(r'[\\/:*?"<>|]')
PLACEHOLDER = re.compile(r"\[[^\]\n]*[A-Za-z][^\]\n]*\]")
LINE_HEIGHT = 13.5      # points per wrapped line at the template's 10 pt Aptos
CHARS_PER_WIDTH = 1.15  # characters that fit per column-width unit at 10 pt
DEFAULT_WIDTH = 8.43
TEMPLATE_SPRINT_SLOTS = 6  # WBS columns E..J
PROGRESS_BY_STATUS = (0, 50, 70, 80, 100, 0)  # template's SWITCH, in the order of its Status drop-down
WBS_BAR_COLOURS = {  # status -> (fill, font), the tones of the template's Quality Metrics rules
    "Done": ("E2F0D9", "375623"),
    "In progress": ("D9EAF7", "17365D"),
    "At risk": ("FFF2CC", "7F6000"),
    "Blocked": ("F4CCCC", "9C0006"),
    "Planned": ("EEF2F7", "243447"),
}

# Template coordinates: the first data row and how many data rows the template ships, per table.
WR_PROGRESS, WR_PROGRESS_ROWS = 14, 6
WR_PEOPLE, WR_PEOPLE_ROWS = 23, 3
WR_NOTE = 27
LEAVE, LEAVE_ROWS = 5, 3
CAPACITY_OFFSET = 6  # first capacity row sits six rows below the leave total
TASKS, TASKS_ROWS = 5, 28
WBS, WBS_ROWS = 5, 15
MILESTONES, MILESTONE_ROWS = 24, 4
METRICS, METRICS_ROWS = 5, 11
METRICS_NOTE_OFFSET = 4  # management note sits four rows below the metrics total
RISKS, RISKS_ROWS = 5, 13
RISKS_NOTE_OFFSET = 2


class DataError(ValueError):
    pass


# --------------------------------------------------------------------------------------------------- data access

def need(obj: dict, key: str, where: str):
    if not isinstance(obj, dict) or obj.get(key) in (None, ""):
        raise DataError(f"{where}.{key} is required")
    return obj[key]


def iso(value, where: str) -> datetime.date:
    try:
        return datetime.date.fromisoformat(value)
    except (TypeError, ValueError):
        raise DataError(f"{where}: expected YYYY-MM-DD, got {value!r}") from None


def dotted(day: datetime.date) -> str:
    return day.strftime("%Y.%m.%d")


def span(period: dict, where: str) -> str:
    return f"{dotted(iso(need(period, 'from', where), where))} ~ {dotted(iso(need(period, 'to', where), where))}"


def allowed_values(ws, column: str) -> tuple:
    """The values the template's own drop-down accepts in `column`, so data is checked against the template."""
    for dv in ws.data_validations.dataValidation:
        if any(get_column_letter(rng.min_col) == column for rng in dv.sqref.ranges):
            return tuple(v.strip() for v in dv.formula1.strip('"').split(","))
    raise DataError(f"template sheet {ws.title!r} has no drop-down in column {column}")


def check(value: str, allowed: tuple, where: str) -> str:
    if value not in allowed:
        raise DataError(f"{where}: {value!r} is not one of {', '.join(allowed)}")
    return value


# --------------------------------------------------------------------------------------------------- layout helpers

def ungroup_columns(ws) -> None:
    """Split grouped column dimensions (e.g. E:J) into one per column so a single column can move or hide."""
    for dim in list(ws.column_dimensions.values()):
        if dim.min is None or dim.max is None or dim.min == dim.max:
            continue
        for index in range(dim.min, dim.max + 1):
            letter = get_column_letter(index)
            ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=dim.width, hidden=dim.hidden)


def width(ws, col: int) -> float:
    letter = get_column_letter(col)
    for dim in ws.column_dimensions.values():
        if dim.min is not None and dim.max is not None and dim.min <= col <= dim.max and dim.width:
            return dim.width
    dim = ws.column_dimensions.get(letter)
    return dim.width if dim is not None and dim.width else DEFAULT_WIDTH


def lines_needed(text, chars: int) -> int:
    return sum(max(1, math.ceil(len(part) / max(chars, 1))) for part in str(text).split("\n"))


def fit_rows(ws, top: int, bottom: int, left: int, right: int, minimum: float) -> None:
    """Size rows top..bottom so the text at (top, left), spanning columns left..right, is not clipped."""
    value = ws.cell(top, left).value
    if value is None or str(value).startswith("="):
        return
    chars = int(sum(width(ws, c) for c in range(left, right + 1)) * CHARS_PER_WIDTH)
    total = max(minimum * (bottom - top + 1), lines_needed(value, chars) * LINE_HEIGHT + 4)
    for row in range(top, bottom + 1):
        ws.row_dimensions[row].height = round(total / (bottom - top + 1), 2)


def fit_data_row(ws, row: int, columns: int, minimum: float, merged: tuple = ()) -> None:
    chars = {c: int(width(ws, c) * CHARS_PER_WIDTH) for c in range(1, columns + 1)}
    for left, right in merged:
        chars[left] = int(sum(width(ws, c) for c in range(left, right + 1)) * CHARS_PER_WIDTH)
    tallest = max(
        (lines_needed(cell.value, chars[cell.column]) for cell in ws[row][:columns]
         if cell.value is not None and not str(cell.value).startswith("=")),
        default=1,
    )
    ws.row_dimensions[row].height = max(minimum or 15.0, tallest * LINE_HEIGHT + 4)


def shift_merges(ws, axis: str, pivot: int, delta: int) -> list:
    """Unmerge every range at or past `pivot` on `axis`; return the shifted or stretched bounds to merge again."""
    pending = []
    for rng in list(ws.merged_cells.ranges):
        lo, hi = (rng.min_row, rng.max_row) if axis == "row" else (rng.min_col, rng.max_col)
        if hi < pivot:
            continue
        bounds = [rng.min_col, rng.min_row, rng.max_col, rng.max_row]
        ws.unmerge_cells(rng.coord)
        if lo >= pivot:
            for index in ((1, 3) if axis == "row" else (0, 2)):
                bounds[index] += delta
        elif delta > 0:  # a banner spanning the insertion point stretches with it
            bounds[3 if axis == "row" else 2] += delta
        else:
            raise DataError(f"{ws.title}: merged range {rng.coord} spans rows that are being removed")
        pending.append(bounds)
    return pending


def merge_again(ws, pending: list) -> None:
    for min_col, min_row, max_col, max_row in pending:
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def retarget(ws, first: int, old_last: int, new_last: int, delta: int) -> None:
    """Point drop-downs and conditional formats that covered the old data block at the new one; shift those below."""
    def moved(rng) -> str:
        left, right = get_column_letter(rng.min_col), get_column_letter(rng.max_col)
        if first <= rng.min_row <= old_last:
            return f"{left}{first}:{right}{new_last}"
        if rng.min_row > old_last:
            return f"{left}{rng.min_row + delta}:{right}{rng.max_row + delta}"
        return rng.coord

    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(" ".join(moved(rng) for rng in dv.sqref.ranges))
    rules = [(cf.sqref, list(cf.rules)) for cf in ws.conditional_formatting]
    ws.conditional_formatting._cf_rules.clear()  # openpyxl has no public way to re-range a rule
    for sqref, cf_rules in rules:
        target = " ".join(moved(rng) for rng in sqref.ranges)
        for rule in cf_rules:
            ws.conditional_formatting.add(target, rule)


def resize_table(ws, first: int, capacity: int, needed: int, columns: int, merged: tuple = ()) -> int:
    """Make the template's data block hold `needed` rows (at least one), restyle it zebra-wise, return its last row.

    `merged` lists (left, right) column spans merged on every data row, e.g. the milestone scope column D:K.
    """
    needed = max(needed, 1)
    delta = needed - capacity
    pivot = first + capacity
    protos = [[copy(ws.cell(first + i, c)._style) for c in range(1, columns + 1)] for i in (0, 1)]
    heights = [ws.row_dimensions[first + i].height for i in (0, 1)]
    for rng in list(ws.merged_cells.ranges):
        if first <= rng.min_row < pivot:
            ws.unmerge_cells(rng.coord)
    if delta:
        pending = shift_merges(ws, "row", pivot, delta)
        below = {r: d.height for r, d in ws.row_dimensions.items() if r >= pivot and d.height}
        for r in below:
            ws.row_dimensions[r].height = None
        if delta > 0:
            ws.insert_rows(pivot, delta)
        else:
            ws.delete_rows(first + needed, -delta)
        for r, h in below.items():
            ws.row_dimensions[r + delta].height = h
        merge_again(ws, pending)
        retarget(ws, first, pivot - 1, first + needed - 1, delta)
    for i in range(needed):
        row = first + i
        for c in range(1, columns + 1):
            cell = ws.cell(row, c)
            cell._style = copy(protos[i % 2][c - 1])
            cell.value = None
        ws.row_dimensions[row].height = heights[i % 2]
        for left, right in merged:
            ws.merge_cells(start_row=row, start_column=left, end_row=row, end_column=right)
    return first + needed - 1


def add_sprint_columns(ws, extra: int) -> None:
    """Insert `extra` sprint columns before Duration, styled like the last template sprint column (J)."""
    pivot = 5 + TEMPLATE_SPRINT_SLOTS
    pending = shift_merges(ws, "col", pivot, extra)
    widths = {c: width(ws, c) for c in range(pivot, ws.max_column + 1)}
    ws.insert_cols(pivot, extra)
    for c, w in sorted(widths.items(), reverse=True):
        letter = get_column_letter(c + extra)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=w)
    for c in range(pivot, pivot + extra):
        letter = get_column_letter(c)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width(ws, pivot - 1))
        for row in range(1, ws.max_row + 1):
            ws.cell(row, c)._style = copy(ws.cell(row, pivot - 1)._style)
    merge_again(ws, pending)


def put(ws, row: int, values: list, start: int = 1) -> None:
    for offset, value in enumerate(values):
        ws.cell(row, start + offset).value = value


def band(ws, coord: str, text: str) -> None:
    """Write a note into a merged band and grow its rows to fit."""
    cell = ws[coord]
    cell.value = text
    for rng in ws.merged_cells.ranges:
        if rng.min_row == cell.row and rng.min_col == cell.column:
            fit_rows(ws, rng.min_row, rng.max_row, rng.min_col, rng.max_col, 15.0)
            return
    fit_rows(ws, cell.row, cell.row, cell.column, cell.column, 15.0)


def leave_days(text: str) -> float:
    """Python mirror of the Leave Plan formula, used only to print the expected value."""
    return text.count("Full day") + 0.5 * (text.count("Morning") + text.count("Afternoon"))


def leave_formula(cell: str) -> str:
    return (f'=(LEN({cell})-LEN(SUBSTITUTE({cell},"Full day","")))/LEN("Full day")'
            f'+0.5*(LEN({cell})-LEN(SUBSTITUTE({cell},"Morning","")))/LEN("Morning")'
            f'+0.5*(LEN({cell})-LEN(SUBSTITUTE({cell},"Afternoon","")))/LEN("Afternoon")')


def breakdown(rng: str, values: tuple, empty: str, sep: str) -> str:
    parts = f'&"{sep}"&'.join(f'"{v} "&COUNTIF({rng},"{v}")' for v in values)
    return f'=IF(COUNTA({rng})=0,"{empty}",{parts})'


# --------------------------------------------------------------------------------------------------- sheets

def fill_tasks(wb, data: dict, anchors: dict) -> dict:
    ws = wb["Task Detail"]
    allowed = allowed_values(ws, "F")
    ws["A2"].value = f"Task detail · {span(data['meta']['period'], 'meta.period')} · update all tasks worked on during the week"
    tasks = data.get("tasks", [])
    last = resize_table(ws, TASKS, TASKS_ROWS, len(tasks), 8)
    progress = ",".join(f'"{s}",{p}%' for s, p in zip(allowed, PROGRESS_BY_STATUS))
    for i, task in enumerate(tasks):
        where, row = f"tasks[{i}]", TASKS + i
        status = check(need(task, "status", where), allowed, f"{where}.status")
        put(ws, row, [i + 1, task.get("storyId") or "—", need(task, "task", where), task.get("module", ""),
                      need(task, "assignee", where), status,
                      f'=IF(F{row}="","",_xlfn.SWITCH(F{row},{progress},0%))', task.get("notes", "")])
        fit_data_row(ws, row, 8, 15.0)
    total, rng = last + 1, f"$F${TASKS}:$F${last}"
    ws[f"F{total}"].value = f'=IF(COUNTA({rng})=0,"",COUNTIF({rng},"Done")&" / "&COUNTA({rng}))'
    ws[f"G{total}"].value = f'=IFERROR(AVERAGE($G${TASKS}:$G${last}),"")'
    ws[f"H{total}"].value = breakdown(rng, allowed, "", "   ·   ")
    ws.auto_filter.ref = f"A4:H{last}"
    ws.freeze_panes = "A5"
    anchors["tasks_total"] = f"F{total}"
    done = sum(1 for t in tasks if t["status"] == "Done")
    return {"tasks complete": f"{done} / {len(tasks)}"}


def fill_leave(wb, data: dict, anchors: dict) -> dict:
    ws = wb["Leave Plan"]
    meta = data["meta"]
    current, following = span(meta["period"], "meta.period"), span(meta["nextPeriod"], "meta.nextPeriod")
    ws["A2"].value = f"Leave plan · current week {current} · next week {following}"
    leave = data.get("leave", [])
    last = resize_table(ws, LEAVE, LEAVE_ROWS, len(leave), 8)
    for i, person in enumerate(leave):
        where, row = f"leave[{i}]", LEAVE + i
        put(ws, row, [i + 1, need(person, "role", where), need(person, "name", where),
                      need(person, "current", where), need(person, "next", where),
                      leave_formula(f"D{row}"), leave_formula(f"E{row}"), person.get("impact", "")])
        fit_data_row(ws, row, 8, 48.0)
    total = last + 1
    ws[f"F{total}"].value = f"=SUM(F{LEAVE}:F{last})"
    ws[f"G{total}"].value = f"=SUM(G{LEAVE}:G{last})"

    capacity = data.get("capacity", [])
    if len(capacity) != 2:
        raise DataError("capacity must have exactly two rows: the current week, then next week")
    first_cap = total + CAPACITY_OFFSET
    for i, week in enumerate(capacity):
        where, row = f"capacity[{i}]", first_cap + i
        days = need(week, "workingDays", where)
        if not isinstance(days, (int, float)) or not 0 <= days <= 7:
            raise DataError(f"{where}.workingDays must be a number of days between 0 and 7")
        put(ws, row, [need(week, "week", where), days, f"=COUNTA($C${LEAVE}:$C${last})", f"=C{row}*D{row}",
                      f"=${'FG'[i]}{total}", f"=E{row}-F{row}", week.get("note", "")], start=2)
        fit_data_row(ws, row, 8, 30.0)
    for col in "EFG":
        ws[f"{col}{first_cap + 2}"].value = f"=SUM({col}{first_cap}:{col}{first_cap + 1})"

    anchors.update(next_available=f"G{first_cap + 1}", next_capacity=f"E{first_cap + 1}")
    days_off = [sum(leave_days(p["current"]) for p in leave), sum(leave_days(p["next"]) for p in leave)]
    next_capacity = capacity[1]["workingDays"] * len(leave)
    return {"leave days (current, next)": days_off,
            "next week available / capacity (md)": f"{next_capacity - days_off[1]} / {next_capacity}"}


def fill_wbs(wb, data: dict) -> None:
    ws = wb["WBS Timeline"]
    wbs = need(data, "wbs", "data")
    statuses = allowed_values(ws, "D")
    sprints = need(wbs, "sprints", "wbs")
    ungroup_columns(ws)
    if len(sprints) > TEMPLATE_SPRINT_SLOTS:
        add_sprint_columns(ws, len(sprints) - TEMPLATE_SPRINT_SLOTS)
    slots = max(len(sprints), TEMPLATE_SPRINT_SLOTS)
    first_col, last_col = 5, 4 + slots
    duration_col = last_col + 1
    dur, last_letter = get_column_letter(duration_col), get_column_letter(last_col)

    starts, ends = [], []
    for i in range(slots):
        letter = get_column_letter(first_col + i)
        if i >= len(sprints):
            ws[f"{letter}4"].value = None
            ws.column_dimensions[letter].hidden = True
            continue
        where = f"wbs.sprints[{i}]"
        start, end = iso(need(sprints[i], "from", where), where), iso(need(sprints[i], "to", where), where)
        starts.append(start)
        ends.append(end)
        ws[f"{letter}4"].value = f"{need(sprints[i], 'name', where)}\n{start:%m.%d} ~ {end:%m.%d}"
    weeks = wbs.get("sprintWeeks", 2)
    caveat = f" · {wbs['caveat']}" if wbs.get("caveat") else ""
    ws["A2"].value = f"WBS timeline · {len(sprints)} sprints × {weeks} weeks · {dotted(min(starts))} → {dotted(max(ends))}{caveat}"

    milestones = wbs.get("milestones", [])
    scope = ((4, duration_col),)
    resize_table(ws, MILESTONES, MILESTONE_ROWS, len(milestones), duration_col, merged=scope)
    for i, item in enumerate(milestones):
        where, row = f"wbs.milestones[{i}]", MILESTONES + i
        put(ws, row, [i + 1, need(item, "name", where), dotted(iso(need(item, "endDate", where), where)),
                      need(item, "scope", where)])
        fit_data_row(ws, row, duration_col, 15.0, merged=scope)

    rows = need(wbs, "rows", "wbs")
    names = [s["name"] for s in sprints]
    last = resize_table(ws, WBS, WBS_ROWS, len(rows), duration_col)
    for i, item in enumerate(rows):
        where, row = f"wbs.rows[{i}]", WBS + i
        put(ws, row, [i + 1, need(item, "phase", where), item.get("owner", ""),
                      check(need(item, "status", where), statuses, f"{where}.status")])
        for name, md in item.get("md", {}).items():
            if name not in names:
                raise DataError(f"{where}.md names sprint {name!r}, which is not in wbs.sprints")
            if not isinstance(md, (int, float)) or md < 0:
                raise DataError(f"{where}.md.{name} must be a non-negative number of man-days")
            ws.cell(row, first_col + names.index(name)).value = md
        ws[f"{dur}{row}"].value = f'=IF(COUNTA(B{row}:{last_letter}{row})=0,"",SUM(E{row}:{last_letter}{row}))'
        fit_data_row(ws, row, 4, 15.0)
    total = last + 1
    for c in range(first_col, duration_col + 1):
        letter = get_column_letter(c)
        ws[f"{letter}{total}"].value = f"=SUM({letter}{WBS}:{letter}{last})"
    bars = f"E{WBS}:{last_letter}{last}"
    for status, (fill, font) in WBS_BAR_COLOURS.items():
        ws.conditional_formatting.add(bars, FormulaRule(
            formula=[f'AND(ISNUMBER(E{WBS}),E{WBS}>0,$D{WBS}="{status}")'],
            fill=PatternFill("solid", start_color=fill, end_color=fill), font=Font(color=font, bold=True)))


def fill_quality(wb, data: dict) -> None:
    ws = wb["Quality Metrics"]
    allowed = allowed_values(ws, "E")
    meta = data["meta"]
    measured = dotted(iso(need(meta, "measuredOn", "meta"), "meta.measuredOn"))
    ws["A2"].value = f"Quality metrics · measured {measured} against {need(meta, 'measuredAgainst', 'meta')}"
    metrics = data.get("metrics", [])
    last = resize_table(ws, METRICS, METRICS_ROWS, len(metrics), 6)
    for i, metric in enumerate(metrics):
        where, row = f"metrics[{i}]", METRICS + i
        put(ws, row, [i + 1, need(metric, "metric", where), need(metric, "result", where), metric.get("target", ""),
                      check(need(metric, "status", where), allowed, f"{where}.status"), metric.get("notes", "")])
        fit_data_row(ws, row, 6, 15.0)
    total, rng = last + 1, f"$E${METRICS}:$E${last}"
    ws[f"E{total}"].value = f'=IF(COUNTA({rng})=0,"—",COUNTIF({rng},"Pass")&"/"&COUNTA({rng}))'
    ws[f"F{total}"].value = breakdown(rng, allowed, "", "   ·   ")
    band(ws, f"A{total + METRICS_NOTE_OFFSET}", f"Management note — {need(data, 'qualityNote', 'data')}")


def fill_risks(wb, data: dict) -> None:
    ws = wb["Risks & Open Questions"]
    allowed = allowed_values(ws, "H")
    risks = data.get("risks", [])
    last = resize_table(ws, RISKS, RISKS_ROWS, len(risks), 8)
    for i, risk in enumerate(risks):
        where, row = f"risks[{i}]", RISKS + i
        put(ws, row, [i + 1, need(risk, "ref", where), need(risk, "item", where), need(risk, "impact", where),
                      risk.get("raised", ""), risk.get("due", ""), risk.get("owner", ""),
                      check(need(risk, "severity", where), allowed, f"{where}.severity")])
        fit_data_row(ws, row, 8, 15.0)
    total, rng = last + 1, f"$H${RISKS}:$H${last}"
    open_levels = tuple(level for level in allowed if level != "Resolved")
    ws[f"F{total}"].value = breakdown(rng, open_levels, "—", "  ·  ")
    ws[f"H{total}"].value = f'=IF(COUNTA({rng})=0,"","✓ "&COUNTIF({rng},"Resolved"))'
    ws.auto_filter.ref = f"A4:H{last}"
    band(ws, f"A{total + RISKS_NOTE_OFFSET}", f"ACTION REQUIRED — {need(data, 'actionRequired', 'data')}")


def fill_weekly(wb, data: dict, anchors: dict) -> None:
    ws = wb["Weekly Report"]
    meta = data["meta"]
    current, following = span(meta["period"], "meta.period"), span(meta["nextPeriod"], "meta.nextPeriod")
    ws["A2"].value = f"Weekly Report · {current} · {need(meta['period'], 'label', 'meta.period')}"
    ws["C4"].value = current
    ws["C5"].value = need(meta, "sprint", "meta")
    ws["C6"].value = need(meta, "overallStatus", "meta")
    ws["C9"].value = need(meta, "preparedBy", "meta")
    ws["C10"].value = dotted(iso(need(meta, "issued", "meta"), "meta.issued"))
    fit_rows(ws, 6, 6, 3, 5, 15.0)
    ws["D13"].value = f"Current week · {current}"
    ws["E13"].value = f"Plan for next week · {following}"
    tasks_total = anchors["tasks_total"]
    ws["C7"].value = f'=IF(\'Task Detail\'!{tasks_total}="","—",\'Task Detail\'!{tasks_total})'
    g, e = anchors["next_available"], anchors["next_capacity"]
    ws["C8"].value = f'=IF(\'Leave Plan\'!{e}=0,"—",\'Leave Plan\'!{g}&" of "&\'Leave Plan\'!{e}&" man-days next week")'

    # Lower table first, so the upper table's template row numbers are still valid when it is resized.
    people = data.get("people", [])
    last = resize_table(ws, WR_PEOPLE, WR_PEOPLE_ROWS, len(people), 5)
    shift = last - (WR_PEOPLE + WR_PEOPLE_ROWS - 1)
    for i, person in enumerate(people):
        where = f"people[{i}]"
        put(ws, WR_PEOPLE + i, [i + 1, need(person, "role", where), need(person, "name", where),
                                need(person, "contribution", where), need(person, "plan", where)])
        fit_data_row(ws, WR_PEOPLE + i, 5, 25.5)

    progress = data.get("progress", [])
    last = resize_table(ws, WR_PROGRESS, WR_PROGRESS_ROWS, len(progress), 5)
    shift += last - (WR_PROGRESS + WR_PROGRESS_ROWS - 1)
    for i, item in enumerate(progress):
        where = f"progress[{i}]"
        put(ws, WR_PROGRESS + i, [i + 1, need(item, "phase", where), need(item, "task", where),
                                  need(item, "current", where), need(item, "next", where)])
        fit_data_row(ws, WR_PROGRESS + i, 5, 15.0)
    band(ws, f"A{WR_NOTE + shift}", f"Capacity note — {need(data, 'capacityNote', 'data')}")


# --------------------------------------------------------------------------------------------------- main

def leftover_placeholders(template: Path, wb) -> list:
    """Cells still holding a placeholder the template ships, e.g. [PROJECT NAME] or [DATE RANGE]."""
    source = load_workbook(template)
    tokens = {m.group(0) for ws in source for row in ws.iter_rows() for cell in row
              if isinstance(cell.value, str) for m in PLACEHOLDER.finditer(cell.value)}
    return [f"{ws.title}!{cell.coordinate}" for ws in wb for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, str) and any(token in cell.value for token in tokens)]


def build(data: dict, template: Path, out_dir: Path) -> tuple:
    meta = need(data, "meta", "data")
    project = need(meta, "project", "meta")
    if FORBIDDEN_IN_NAME.search(project):
        raise DataError(f"meta.project {project!r} cannot be used in a file name")
    start = iso(need(need(meta, "period", "meta"), "from", "meta.period"), "meta.period.from")
    end = iso(need(meta["period"], "to", "meta.period"), "meta.period.to")
    need(meta, "nextPeriod", "meta")

    wb = load_workbook(template)
    for ws in wb:
        ws["A1"].value = f"{project} — PROJECT WEEKLY REPORT"
    anchors: dict = {}
    expected = {}
    expected.update(fill_tasks(wb, data, anchors))
    expected.update(fill_leave(wb, data, anchors))
    fill_wbs(wb, data)
    fill_quality(wb, data)
    fill_risks(wb, data)
    fill_weekly(wb, data, anchors)  # last: it points at cells the other sheets placed
    left = leftover_placeholders(template, wb)
    if left:
        raise DataError(f"template placeholders left unfilled: {', '.join(left)}")

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{project}_Weekly Report_{start:%Y%m%d}~{end:%Y%m%d}.xlsx"
    wb.save(path)
    return path, expected


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill the ADF weekly-report workbook template from a JSON data file.")
    parser.add_argument("--data", required=True, type=Path, help="weekly data JSON (references/weekly-data-contract.md)")
    parser.add_argument("--out", required=True, type=Path, help="directory to write the workbook into")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="workbook template")
    args = parser.parse_args()
    if not args.template.is_file():
        print(f"template not found: {args.template}", file=sys.stderr)
        return 2
    try:
        data = json.loads(args.data.read_text(encoding="utf-8"))
        path, expected = build(data, args.template, args.out)
    except (OSError, json.JSONDecodeError, DataError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    print(f"written: {path}")
    for label, value in expected.items():
        print(f"expected {label}: {value}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
